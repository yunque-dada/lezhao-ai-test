# -*- coding: utf-8 -*-
"""生成Python海龟画图课程PPT"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os
import sys

# 设置输出编码
sys.stdout.reconfigure(encoding='utf-8')

# 创建PPT
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 定义颜色
PRIMARY = RGBColor(99, 102, 241)
ACCENT = RGBColor(236, 72, 153)
DARK = RGBColor(30, 27, 75)
LIGHT = RGBColor(248, 250, 252)

def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 背景
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK
    shape.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = LIGHT
    p.alignment = PP_ALIGN.CENTER
    
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12.333), Inches(1.5))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(24)
        p.font.color.rgb = LIGHT
        p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, content_items, code=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 左侧装饰条
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.15), prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PRIMARY
    shape.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    
    # 内容
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(6), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, item in enumerate(content_items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + item if item else ""
        p.font.size = Pt(22)
        p.font.color.rgb = DARK
        p.space_after = Pt(12)
    
    # 代码框
    if code:
        # 代码背景
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.2), Inches(6), Inches(5.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(242, 242, 247)
        shape.line.fill.background()
        
        code_box = slide.shapes.add_textbox(Inches(7), Inches(1.3), Inches(5.8), Inches(5.5))
        tf = code_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = code
        p.font.size = Pt(13)
        p.font.name = 'Consolas'
        p.font.color.rgb = RGBColor(15, 23, 42)
    
    return slide

# ===== 第1页：封面 =====
add_title_slide(prs, "神奇的万花筒", "用Python画出炫彩图案\nPython海龟画图进阶课程\n适合年级：小学3～6年级")

# ===== 第2页：课程目标 =====
add_content_slide(prs, "今天我们要学什么？", [
    "回顾海龟画图的基本命令",
    "学会用for循环重复画图",
    "学会用random模块让颜色随机变化",
    "学会用旋转和大小变化创造炫酷效果",
    "自己动手设计独一无二的万花筒图案！"
])

# ===== 第3页：导入 =====
add_content_slide(prs, "你玩过万花筒吗？", [
    "万花筒是一种光学玩具，里面有多彩的图案",
    "转动万花筒，图案会不断变化，非常漂亮",
    "",
    "今天我们就要用Python来创造数字万花筒！"
])

# ===== 第4页：准备工作 =====
code4 = '''import turtle
import random

t = turtle.Turtle()
t.speed(0)
turtle.bgcolor("black")
t.pensize(2)'''
add_content_slide(prs, "让我们开始吧！", ["打开Python（IDLE或Thonny）", "导入需要的模块：turtle 和 random", "设置画笔速度和背景颜色"], code=code4)

# ===== 第5页：颜色列表 =====
code5 = '''colors = ["red", "orange", "yellow", 
          "green", "blue", "purple", 
          "pink", "cyan", "magenta"]'''
add_content_slide(prs, "五颜六色的画笔", ["创建一个颜色列表", "让程序随机挑选颜色", "颜色名字用英文"], code=code5)

# ===== 第6页：正八边形 =====
code6 = '''for side in range(8):
    t.forward(100)
    t.right(45)'''
add_content_slide(prs, "先画一个简单的图形", ["正八边形有8条边", "每次右转45度", "边长100像素"], code=code6)

# ===== 第7页：颜色随机 =====
code7 = '''for side in range(8):
    t.color(random.choice(colors))
    t.forward(100)
    t.right(45)'''
add_content_slide(prs, "给图形穿上彩色外衣", ["每次画边之前", "从颜色列表中随机选一个颜色"], code=code7)

# ===== 第8页：旋转 =====
code8 = '''for times in range(36):
    for side in range(8):
        t.color(random.choice(colors))
        t.forward(100)
        t.right(45)
    t.right(10)'''
add_content_slide(prs, "转起来！变成万花筒", ["画完一个八边形后", "将海龟旋转一个小角度（比如10度）", "再画下一个八边形", "形成旋转效果"], code=code8)

# ===== 第9页：改变大小 =====
code9 = '''size = 50
for times in range(30):
    for side in range(8):
        t.color(random.choice(colors))
        t.forward(size)
        t.right(45)
    t.right(15)
    size = size + 3'''
add_content_slide(prs, "让图案有层次感", ["每次画完一圈", "把边长增加一点点", "形成螺旋上升的效果"], code=code9)

# ===== 第10页：完整代码 =====
code10 = '''import turtle
import random

t = turtle.Turtle()
t.speed(0)
turtle.bgcolor("black")
t.pensize(2)

colors = ["red", "orange", "yellow", 
          "green", "blue", "purple", 
          "pink", "cyan", "magenta"]

size = 20
for times in range(72):
    t.color(random.choice(colors))
    for side in range(30):
        t.forward(size)
        t.right(91)
        size += 1
    t.right(5)

t.hideturtle()
turtle.done()'''
add_content_slide(prs, "一起来运行吧！", ["这是我们这节课的最终代码", "可以直接复制运行"], code=code10)

# ===== 第11页：互动活动 =====
add_content_slide(prs, "变身小小艺术家", [
    "外循环次数：72 -> 36 或 120",
    "内循环次数：30 -> 10 或 50",
    "旋转角度：5 -> 10 或 2",
    "颜色列表：添加自己喜欢的颜色",
    "",
    "和同桌比一比，谁画出的图案最漂亮！"
])

# ===== 第12页：作品展示 =====
add_content_slide(prs, "展示你的万花筒", [
    "运行你的程序，保存截图",
    "",
    "向大家介绍你的作品：",
    "用了哪些颜色？改了哪些参数？",
    "",
    "说说你最喜欢的地方"
])

# ===== 第13页：课程总结 =====
add_content_slide(prs, "今天我们学会了……", [
    "用循环画出重复图案",
    "用随机数让颜色多变",
    "通过旋转和大小变化创造动态效果",
    "",
    "编程不仅能解决问题，还能创作艺术！"
])

# ===== 第14页：课后作业 =====
add_content_slide(prs, "课后可以继续探索", [
    "颜色渐变：试试不用随机，让颜色按顺序变化",
    "",
    "图形变形：把正八边形改成五边形或圆形",
    "",
    "家庭作业：创作「送给妈妈的万花筒」"
])

# ===== 第15页：结束页 =====
add_title_slide(prs, "谢谢大家！", "期待你们的创意作品\nQ&A环节")

# 保存PPT
output_path = os.path.join(os.path.dirname(__file__), "turtle-course.pptx")
prs.save(output_path)
print("PPT saved to: " + output_path)
